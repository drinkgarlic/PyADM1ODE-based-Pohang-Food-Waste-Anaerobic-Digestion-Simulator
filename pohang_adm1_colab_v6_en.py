# ============================================================
# Colab: ZIP -> PyADM1ODE -> monthly YAML -> dynamic gas rates
# v6-EN (2026-10-14) — English edition of v6 (identical logic; comments, messages and labels in English)
#   - Dedicated H2S / NH3 concentration plots in ppm (ADM1_H2S_NH3_ppm.png: two panels, monthly mean and max)
#   - Feedstock input: reads sheet '표2_YAML변수' (17 variables x 12 months) of /content/pohang_fw_yaml_monthly.xlsx
#       and writes data/substrates/pohang_waste_MM.yaml (12 files) automatically — no hand-made YAML needed
#   - Dilution: YAML holds 1:9-diluted concentrations (intensive properties); flow fixed at Q_FEED = 200 m3/d
#       (Q_feed in the Excel 'assumptions' sheet)
#   - Warm-up phase: November feed, Q ramped in steps (not recorded)
#   - Gas composition (CH4/CO2/H2, dry and wet basis) + moisture (saturated water vapour) columns
#   - 1-second time series: hourly ODE results linearly interpolated to 1 s (parquet/zstd; csv.gz if pyarrow is missing)
#   - K_I_NH3_FACTOR: scaling factor for the free-ammonia inhibition constant (model parameter)
#   - Ammonia: equilibrium NH3 in off-gas (ppmv, g/d, Henry's law) + ammonium nitrogen leaving in digestate (kg N/d)
#   - Hydrogen sulphide: protein sulphur -> sulphide production (using the model's protein degradation extent)
#       -> pH/temperature partitioning (dissociation + Henry's law) -> H2S in off-gas.
#       ADM1 has no sulphur cycle, so this is an equilibrium (upper-bound) estimate. Assumptions: F_S_PROTEIN, SO4_IN_mgS_L, FE_DOSE_mol_m3
# Set RUN_MODE='test' first, then change to 'full'.
# ============================================================
from google.colab import drive
from pathlib import Path
from zipfile import ZipFile
import sys, os, shutil, time, warnings, signal, gzip
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import yaml

# ---------------- User settings ----------------
RUN_MODE = "full"         # "test" = 2 days, "full" = 12 months
STEP_HOURS = 1            # ADM1 integration step: 1, 2, 3, 4, 6, 12, 24
WRITE_SECONDS = True      # True: write 1-second time series (one file per month)
SECONDS_FORMAT = "parquet"  # "parquet" (needs pyarrow; falls back to csv.gz automatically) | "csv"
YAML_DIR_OVERRIDE = None  # Optional alternative YAML folder (None -> ROOT/data/substrates)
START_MONTH = "2020-11"   # November through next October

XLSX_PATH = "/content/pohang_fw_yaml_monthly.xlsx"   # feedstock workbook (if absent, a file of the same name in Drive root is used)
XLSX_SHEET = "표2_YAML변수"                           # sheet name as it exists in the workbook ("Table 2 YAML variables"): 17 variables x 12 months
YAML_PREFIX = "pohang_waste"                          # generated file names: pohang_waste_01.yaml ... _12.yaml
REBUILD_YAML = True                                   # False: skip the Excel and use the existing YAML files

Q_FEED = 200.0            # m3/d  slurry flow into the digester (Q_feed in the Excel assumptions sheet). YAML concentrations are 1:9-diluted values
W_WATER = 9.0             # t water / t raw waste (Excel w_water) — for reference: raw waste in this flow = Q_FEED/(1+W_WATER) = 20 t/d
V_LIQ = 6000.0            # m3  digester working volume (one tank)
V_GAS = 800.0             # m3  headspace
T_AD = 308.15             # K = 35 C

WARMUP_DAYS = 30                       # days per warm-up step (0 = skip warm-up)
WARMUP_FRACTIONS = (0.25, 0.50, 0.75, 1.00)
K_I_NH3_FACTOR = 1.0      # ADM1da default K_I,NH3 = 0.0018 kmol/m3 (= 25 mg N/L); 4.0 -> 100 mg N/L

# --- Hydrogen sulphide estimate: assumptions (computed outside ADM1) ---
F_S_PROTEIN = 0.010       # sulphur mass fraction of protein [kg S / kg protein]; food proteins 0.8-1.2 % (Cys, Met)
SO4_IN_mgS_L = 0.0        # sulphate in the feed slurry [mg S/L] — 0 if unknown (assumed fully reduced to sulphide)
FE_DOSE_mol_m3 = 0.0      # iron salt dose [mol Fe / m3 slurry] — assumed to precipitate sulphide 1:1 as FeS (0 = none)
# ------------------------------------------------

assert RUN_MODE in ("test", "full")
assert STEP_HOURS > 0 and 24 % STEP_HOURS == 0

# ============================================================
# 1. Mount and extract ZIP only when necessary
# ============================================================

if not Path("/content/drive/MyDrive").exists():
    drive.mount("/content/drive")

ZIP_PATH = Path("/content/drive/MyDrive/PyADM1ODE-master.zip")
EXTRACT_DIR = Path("/content/PyADM1ODE")

ROOT_OPTIONS = [
    EXTRACT_DIR / "PyADM1ODE-master",
    EXTRACT_DIR,
    Path("/content/PyADM1ODE-master"),
]

ROOT = next((p for p in ROOT_OPTIONS if (p / "pyadm1" / "__init__.py").is_file()), None)

if ROOT is None:
    if not ZIP_PATH.exists():
        raise FileNotFoundError(f"ZIP missing: {ZIP_PATH}")
    print("Extracting ZIP into Colab local disk...", flush=True)
    EXTRACT_DIR.mkdir(parents=True, exist_ok=True)
    with ZipFile(ZIP_PATH) as zf:
        zf.extractall(EXTRACT_DIR)
    ROOT = next((p for p in ROOT_OPTIONS if (p / "pyadm1" / "__init__.py").is_file()), None)

if ROOT is None:
    raise FileNotFoundError("pyadm1/__init__.py not found after extraction. Check ZIP directory structure.")

os.chdir(ROOT)
sys.path.insert(0, str(ROOT))
print("Project root:", ROOT, flush=True)

if "pyadm1" in sys.modules:
    previous = getattr(sys.modules["pyadm1"], "__file__", "") or ""
    if not str(Path(previous).resolve()).startswith(str(ROOT.resolve())):
        for name in list(sys.modules):
            if name == "pyadm1" or name.startswith("pyadm1."):
                del sys.modules[name]

# ============================================================
# 2. Build monthly YAML files from the Excel sheet, then check them
# ============================================================

SUBSTRATES = Path(YAML_DIR_OVERRIDE) if YAML_DIR_OVERRIDE else ROOT / "data" / "substrates"
SUBSTRATES.mkdir(parents=True, exist_ok=True)

YAML_KEYS = ["TS", "NH4", "BGP", "BMP", "aXI", "fOTSrf", "fsOTS", "ffOTS", "aSi",
             "fRF", "fRP", "fRFe", "fRA", "Temp", "pH", "KS43", "FFS"]          # order of PyADM1ODE SubstrateParams
YAML_UNITS = {"TS": "kg/t FM", "NH4": "kg N/m3", "BGP": "Nm3/t VS", "BMP": "Nm3 CH4/t VS", "aXI": "-", "fOTSrf": "-",
              "fsOTS": "-", "ffOTS": "-", "aSi": "-", "fRF": "-TS", "fRP": "-TS", "fRFe": "-TS", "fRA": "-TS",
              "Temp": "C", "pH": "-", "KS43": "mol/m3", "FFS": "kg HAc/m3"}

def read_table2(xlsx_path, sheet):
    """Read the 17 variables x 12 months from the Excel sheet and return {month: {var: value}}."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)   # data_only: read the cached results of formulas
    if sheet not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet}' not found. Available sheets: {wb.sheetnames}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    # Header row: column A reads '변수' (= 'variable'); the same row holds the month headers '1월' ... '12월' (= 'Jan' ... 'Dec')
    hdr_idx = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == "변수"), None)
    if hdr_idx is None:
        raise ValueError("Header row (column A == '변수') not found in the sheet.")
    hdr = rows[hdr_idx]
    month_col = {}
    for j, v in enumerate(hdr):
        if isinstance(v, str) and v.strip().endswith("월"):      # '월' = month suffix in the Korean header
            try:
                month_col[int(v.strip()[:-1])] = j
            except ValueError:
                pass
    missing_m = [m for m in range(1, 13) if m not in month_col]
    if missing_m:
        raise ValueError(f"Month columns missing in the header row: {missing_m}")
    var_row = {}
    for r in rows[hdr_idx + 1:]:
        if r and isinstance(r[0], str) and r[0].strip() in YAML_KEYS and r[0].strip() not in var_row:
            var_row[r[0].strip()] = r
        if len(var_row) == len(YAML_KEYS):
            break
    missing_v = [k for k in YAML_KEYS if k not in var_row]
    if missing_v:
        raise ValueError(f"Variables not found in the sheet: {missing_v}")
    data = {m: {} for m in range(1, 13)}
    for k in YAML_KEYS:
        for m in range(1, 13):
            v = var_row[k][month_col[m]]
            if v is None or not isinstance(v, (int, float)):
                raise ValueError(
                    f"Empty value for {k}, month {m} ({v!r}). If the workbook has no cached formula results, open and save it in Excel (recalculate), then upload again.")
            data[m][k] = float(v)
    # Cross-check: Q_feed in the assumptions sheet ('가정') vs. Q_FEED in this notebook (warning only)
    if "가정" in wb.sheetnames:
        for r in wb["가정"].iter_rows(values_only=True):
            if r and r[0] == "Q_feed" and isinstance(r[2], (int, float)) and abs(float(r[2]) - Q_FEED) > 1e-6:
                warnings.warn(f"Q_feed in the Excel assumptions sheet ({r[2]}) differs from Q_FEED in the notebook ({Q_FEED}).")
    wb.close()
    return data

def write_month_yaml(month, vals, path, source):
    lines = [f"# ADM1da substrate definition: Pohang food waste, month {month:02d}",
             f"# generated from {source} / sheet '{XLSX_SHEET}' on {time.strftime('%Y-%m-%d %H:%M')}",
             "# Model-default constants (densities, COD factors, acid-base constants) are inherited from SubstrateParams.",
             "",
             f"name: Pohang food waste {month:02d}", ""]
    for k in YAML_KEYS:
        lines.append(f"{k}: {vals[k]:.6g}    # {YAML_UNITS[k]}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

monthly_files = {m: SUBSTRATES / f"{YAML_PREFIX}_{m:02d}.yaml" for m in range(1, 13)}

if REBUILD_YAML:
    xlsx = Path(XLSX_PATH)
    if not xlsx.is_file():
        alt = Path("/content/drive/MyDrive") / xlsx.name
        if alt.is_file():
            xlsx = alt
        else:
            raise FileNotFoundError(f"Excel file not found: {XLSX_PATH} (also not in Drive root). Upload it via the Colab file panel.")
    table2 = read_table2(xlsx, XLSX_SHEET)
    for m in range(1, 13):
        write_month_yaml(m, table2[m], monthly_files[m], xlsx.name)
    print(f"12 monthly YAML files written to {SUBSTRATES}  (source: {xlsx})", flush=True)
    chk = pd.DataFrame(table2).T[YAML_KEYS]
    chk.index.name = "month"
    print(chk.round(4).to_string(), flush=True)

for m in range(1, 13):
    p = monthly_files[m]
    if not p.is_file():
        raise FileNotFoundError(f"Missing: {p}\nSet REBUILD_YAML=True to generate it from the Excel file, or change YAML_DIR_OVERRIDE.")
    with p.open(encoding="utf-8") as fp:
        data = yaml.safe_load(fp)
    missing_k = [k for k in YAML_KEYS if k not in data]
    if missing_k:
        raise ValueError(f"{p.name}: missing keys {missing_k}")
    if float(data["KS43"]) < 1.0:
        warnings.warn(f"{p.name}: KS43={data['KS43']} is low. Verify alkalinity before QRA.")
    if abs(float(data["fsOTS"]) + float(data["ffOTS"]) - 1.0) > 1e-3:
        warnings.warn(f"{p.name}: fsOTS+ffOTS={float(data['fsOTS'])+float(data['ffOTS']):.3f} (should be 1)")
print("All 12 monthly YAML files ready.", flush=True)

# ============================================================
# 3. Local output first, Google Drive backup second
# ============================================================

OUT_LOCAL = Path("/content/ADM1_QRA_results") / RUN_MODE
OUT_DRIVE = Path("/content/drive/MyDrive/ADM1_QRA_results") / RUN_MODE
OUT_LOCAL.mkdir(parents=True, exist_ok=True)
OUT_DRIVE.mkdir(parents=True, exist_ok=True)

def save_copy(path):
    shutil.copy2(path, OUT_DRIVE / path.name)

# ============================================================
# 4. Import PyADM1ODE
# ============================================================

from pyadm1.configurator.plant_builder import BiogasPlant
from pyadm1.configurator.plant_configurator import PlantConfigurator
from pyadm1.substrates.feedstock import Feedstock
import pyadm1
print("Loaded library:", pyadm1.__file__, flush=True)

def feedstock_for(month):
    return Feedstock(str(monthly_files[month]), feeding_freq=STEP_HOURS, total_simtime=600, simba_q_convention=False)

# ============================================================
# 5. Initialize plant ONCE, apply K_I override, run warm-up
# ============================================================

start_period = pd.Period(START_MONTH, freq="M")
periods = [start_period + i for i in range(1 if RUN_MODE == "test" else 12)]

first_feed = feedstock_for(periods[0].month)
plant = BiogasPlant("Pohang_QRA")
config = PlantConfigurator(plant, first_feed)
digester, state_info = config.add_digester(
    digester_id="main_digester", V_liq=V_LIQ, V_gas=V_GAS, T_ad=T_AD, Q_substrates=[Q_FEED])
plant.initialize()
print("Initialization:", str(state_info).strip(), flush=True)

# --- Model parameter adjustment (independent of the feedstock YAML) ---
_ip = digester.adm1._inhib_params
_ip["K_I_nh3"] *= K_I_NH3_FACTOR
_ip["K_I_nh3_pro"] *= K_I_NH3_FACTOR
print(f"K_I_nh3 = {_ip['K_I_nh3']:.5f} kmol/m3 (= {_ip['K_I_nh3']*14000:.0f} mg N/L)", flush=True)

# --- Gas-phase constants (taken from the library itself) ---
ADM = digester.adm1
P_H2O_BAR = ADM._p_gas_h2o                     # saturated water vapour pressure at 35 C ~ 0.0557 bar
RHO_W_NORM_G = 1.01325e5 * 0.018015 / (8.314 * 293.15) * 1000.0   # density of pure water vapour at 20 C, 1.01325 bar ~ 749 g/m3
RHO_W_ACT_G = P_H2O_BAR * 1e5 * 0.018015 / (8.314 * T_AD) * 1000.0  # saturated absolute humidity at digester temperature ~ 39 g/m3
_IDX_P_H2, _IDX_P_CH4, _IDX_P_CO2, _IDX_P_TOT = 37, 38, 39, 40
_IDX_S_AC, _IDX_S_IN, _IDX_S_NH3 = 6, 10, 36

# --- Ammonia in off-gas (equilibrium estimate): ADM1 tracks NH3 only in the liquid (S_IN, S_nh3); there is no gas-phase NH3.
#     Henry's law gives the equilibrium NH3 concentration of the off-gas (upper bound). Sander (2015): H_cp(298 K) = 0.59 mol/(m3*Pa), -dlnH/d(1/T) = 4200 K
H_NH3 = 0.59 * np.exp(4200.0 * (1.0 / T_AD - 1.0 / 298.15))        # mol/(m3·Pa) @ T_AD
RHO_NH3_NORM_G = 1.01325e5 * 0.017031 / (8.314 * 293.15) * 1000.0  # density of pure NH3 at 20 C ~ 708 g/m3

# --- H2S partitioning constants: Sander (2015) H_cp(H2S, 298 K) = 1.0e-3 mol/(m3*Pa), -dlnH/d(1/T) = 2100 K; pKa1(25 C) = 7.05, dH ~ 22 kJ/mol
H_H2S = 1.0e-3 * np.exp(2100.0 * (1.0 / T_AD - 1.0 / 298.15))
PKA_H2S = 7.05 - (22000.0 / (2.303 * 8.314)) * (1.0 / 298.15 - 1.0 / T_AD)
KA_H2S = 10.0 ** (-PKA_H2S)
RHO_H2S_NORM_G = 1.01325e5 * 0.034081 / (8.314 * 293.15) * 1000.0   # ~ 1417 g/m3
RT_NORM = 8.314 * 293.15
_IDX_S_AA, _IDX_X_PS_PR, _IDX_X_PF_PR, _IDX_X_S_PR = 1, 13, 16, 19
_protein_in = {"PF": 0.0, "PS": 0.0}     # degradable protein COD of the current feed (updated after every set_feedstock)

def update_protein_in(Q):
    c = digester.feedstock.blended_concentrations([Q])
    _protein_in["PF"] = float(c.get("X_PF_pr", 0.0)); _protein_in["PS"] = float(c.get("X_PS_pr", 0.0))
print(f"H_H2S = {H_H2S:.2e} mol/(m3·Pa), pKa1(H2S) = {PKA_H2S:.2f} @ {T_AD-273.15:.0f} C", flush=True)

def gas_record(result):
    """Result dict + ADM1 state vector -> gas flow, composition and moisture columns (gas at 20 C, 1.01325 bar; Q_gas is wet)"""
    s = np.asarray(digester.adm1_state, dtype=float)
    q_gas, q_ch4, q_co2, q_h2o, p_dry = ADM.calc_gas(s[_IDX_P_H2], s[_IDX_P_CH4], s[_IDX_P_CO2], s[_IDX_P_TOT])
    p_wet = p_dry + P_H2O_BAR
    q_h2 = q_gas * s[_IDX_P_H2] / p_wet if p_wet > 0 else 0.0
    q_dry = q_gas - q_h2o
    dry = lambda p: 100.0 * p / p_dry if p_dry > 0 else np.nan
    out = {
        "Q_gas_wet_m3_d_20C": q_gas,
        "Q_gas_dry_m3_d_20C": q_dry,
        "Q_CH4_m3_d": q_ch4,
        "Q_CO2_m3_d": q_co2,
        "Q_H2_m3_d": q_h2,
        "Q_H2O_vapour_m3_d": q_h2o,
        "CH4_dry_pct": dry(s[_IDX_P_CH4]),
        "CO2_dry_pct": dry(s[_IDX_P_CO2]),
        "H2_dry_ppm": 1e4 * dry(s[_IDX_P_H2]),
        "CH4_wet_pct": 100.0 * s[_IDX_P_CH4] / p_wet,
        "CO2_wet_pct": 100.0 * s[_IDX_P_CO2] / p_wet,
        "H2O_wet_pct": 100.0 * P_H2O_BAR / p_wet,
        "H2O_g_per_m3_wet_20C": P_H2O_BAR / p_wet * RHO_W_NORM_G,   # water per m3 of wet gas (20 C basis)
        "H2O_g_per_m3_at_35C": RHO_W_ACT_G,                           # saturated absolute humidity at digester temperature (dew point = 35 C)
        "H2O_kg_d": q_h2o * RHO_W_NORM_G / 1000.0,
        "p_gas_dry_bar": p_dry,
        "p_gas_wet_bar": p_wet,
        "pH": float(result.get("pH", np.nan)),
        "VFA_kgCOD_m3": float(result.get("VFA", np.nan)),
        "S_ac_kgCOD_m3": s[_IDX_S_AC],
        "TAN_gN_L": s[_IDX_S_IN] * 14.0,
        "NH3_mgN_L": s[_IDX_S_NH3] * 14000.0,
        # Ammonia (estimate and balance)
        "NH3_gas_ppmv_eq": s[_IDX_S_NH3] * 1000.0 / H_NH3 / 1e5 / 1.013 * 1e6,   # kmol -> mol/m3 -> bar -> ppmv
        "NH3_gas_g_d_eq": q_dry * (s[_IDX_S_NH3] * 1000.0 / H_NH3 / 1e5 / 1.013) * RHO_NH3_NORM_G,
        "TAN_out_kgN_d": Q_FEED * s[_IDX_S_IN] * 14.0,                              # ammonium nitrogen leaving with the digestate
    }
    # --- Hydrogen sulphide (equilibrium estimate) ---
    pr_in = _protein_in["PF"] + _protein_in["PS"]
    pr_res = s[_IDX_X_PS_PR] + s[_IDX_X_PF_PR] + s[_IDX_X_S_PR] + s[_IDX_S_AA]
    eta_pr = max(0.0, 1.0 - pr_res / pr_in) if pr_in > 0 else 0.0           # fraction of degradable protein actually degraded
    S_prod = Q_FEED * pr_in * eta_pr * F_S_PROTEIN * 0.7736 / 0.032 + Q_FEED * SO4_IN_mgS_L / 1000.0 / 32.06   # mol S/d
    S_eff = max(0.0, S_prod - Q_FEED * FE_DOSE_mol_m3)
    S_H = 10.0 ** (-out["pH"])
    a0 = S_H / (S_H + KA_H2S)                                                # fraction present as H2S(aq)
    gterm = q_dry * a0 / (H_H2S * RT_NORM)
    C_tot = S_eff / (Q_FEED + gterm) if (Q_FEED + gterm) > 0 else 0.0       # total dissolved sulphide in the liquid [mol S/m3]
    p_h2s = a0 * C_tot / H_H2S                                               # Pa
    out.update({
        "protein_deg_frac": eta_pr,
        "S_released_kgS_d": S_eff * 0.032,
        "H2S_gas_ppmv_eq": p_h2s / 1.01325e5 * 1e6,
        "H2S_gas_kg_d_eq": q_dry * p_h2s / 1.01325e5 * RHO_H2S_NORM_G / 1000.0,
        "sulfide_liq_mgS_L": C_tot * 32.06,
        "sulfide_out_kgS_d": Q_FEED * C_tot * 0.032,
    })
    return out

dt_day = STEP_HOURS / 24.0
steps_per_day = 24 // STEP_HOURS

if RUN_MODE == "full" and WARMUP_DAYS > 0:
    for frac in WARMUP_FRACTIONS:
        digester.set_feedstock(feedstock_for(periods[0].month), Q_substrates=[Q_FEED * frac], rebuild_state=False)
        update_protein_in(Q_FEED * frac)
        for _ in range(WARMUP_DAYS * steps_per_day):
            _r = plant.step(dt_day)["main_digester"]
        g = gas_record(_r)
        print(f"warm-up Q={Q_FEED*frac:6.0f} m3/d x {WARMUP_DAYS} d: pH={g['pH']:.2f} VFA={g['VFA_kgCOD_m3']:.2f} "
              f"CH4={g['Q_CH4_m3_d']:.0f} m3/d CH4dry={g['CH4_dry_pct']:.1f}%", flush=True)
    digester.set_feedstock(feedstock_for(periods[0].month), Q_substrates=[Q_FEED], rebuild_state=False)
update_protein_in(Q_FEED)

# ============================================================
# 6. Diagnostic timeout for individual ODE steps
# ============================================================

MAX_STEP_SECONDS = 60 if RUN_MODE == "test" else 120

def on_step_timeout(signum, frame):
    raise TimeoutError(f"An ADM1 BDF step exceeded {MAX_STEP_SECONDS} seconds. Check pH, KS43, initial state and solver settings.")

signal.signal(signal.SIGALRM, on_step_timeout)

# ============================================================
# 7. Continuous simulation with progress logging
# ============================================================

SEC_COLS = ["Q_gas_wet_m3_d_20C", "Q_gas_dry_m3_d_20C", "Q_CH4_m3_d", "Q_CO2_m3_d", "Q_H2_m3_d", "Q_H2O_vapour_m3_d",
            "CH4_dry_pct", "CO2_dry_pct", "H2_dry_ppm", "CH4_wet_pct", "CO2_wet_pct", "H2O_wet_pct",
            "H2O_g_per_m3_wet_20C", "NH3_gas_ppmv_eq", "H2S_gas_ppmv_eq", "pH", "VFA_kgCOD_m3"]

all_frames = []
last_record = None          # last value of the previous phase (first knot of the 1-second interpolation)

for month_index, period in enumerate(periods):

    if month_index > 0:
        old_state = np.asarray(digester.adm1_state, dtype=float).copy()
        digester.set_feedstock(feedstock_for(period.month), Q_substrates=[Q_FEED], rebuild_state=False)
        update_protein_in(Q_FEED)
        new_state = np.asarray(digester.adm1_state, dtype=float)
        if not np.array_equal(old_state, new_state, equal_nan=True):
            raise RuntimeError("Biological state reset at monthly boundary.")

    days = 2 if RUN_MODE == "test" else period.days_in_month
    n_steps = days * steps_per_day
    month_start = period.start_time
    rows = []

    print("\n" + "=" * 60, flush=True)
    print(f"SIMULATING {period}: {days} days, {n_steps} steps", flush=True)
    print("=" * 60, flush=True)
    month_clock = time.perf_counter()

    for i in range(n_steps):
        if i == 0 or (i + 1) % steps_per_day == 0:
            print(f"Calling BDF step {i+1}/{n_steps}...", flush=True)
        t0 = time.perf_counter()
        signal.setitimer(signal.ITIMER_REAL, MAX_STEP_SECONDS)
        try:
            result = plant.step(dt_day)["main_digester"]
        finally:
            signal.setitimer(signal.ITIMER_REAL, 0)
        elapsed = time.perf_counter() - t0
        if elapsed > 20:
            print(f"SLOW STEP: {elapsed:.1f} seconds at step {i+1}", flush=True)

        g = gas_record(result)
        if not np.all(np.isfinite([g["Q_gas_wet_m3_d_20C"], g["Q_CH4_m3_d"], g["Q_CO2_m3_d"]])) \
           or min(g["Q_gas_wet_m3_d_20C"], g["Q_CH4_m3_d"], g["Q_CO2_m3_d"]) < 0:
            raise RuntimeError(f"Invalid gas result: {period}, step {i+1}")

        interval_start = month_start + pd.Timedelta(hours=i * STEP_HOURS)
        row = {"interval_start": interval_start,
               "interval_end": interval_start + pd.Timedelta(hours=STEP_HOURS),
               "month": str(period), "Q_in_m3_d": Q_FEED}
        row.update(g)
        row["integration_sec"] = elapsed
        rows.append(row)

    month_df = pd.DataFrame(rows)
    fp = OUT_LOCAL / f"ADM1_{period}_steps.csv"
    month_df.to_csv(fp, index=False)
    save_copy(fp)
    all_frames.append(month_df)

    print(f"Completed {period} in {time.perf_counter()-month_clock:.1f} seconds", flush=True)
    print(f"Last pH = {g['pH']:.3f}; CH4(dry) = {g['CH4_dry_pct']:.1f}%; H2(dry) = {g['H2_dry_ppm']:.0f} ppm; "
          f"H2O = {g['H2O_wet_pct']:.2f}% ({g['H2O_g_per_m3_wet_20C']:.1f} g/m3); "
          f"NH3(eq) = {g['NH3_gas_ppmv_eq']:.0f} ppm; H2S(eq) = {g['H2S_gas_ppmv_eq']:.0f} ppm", flush=True)

    # --------------------------------------------------------
    # 1-second time series: linear interpolation between interval-end knots, written day by day to limit memory
    # --------------------------------------------------------
    if WRITE_SECONDS:
        knot_prev = last_record if last_record is not None else month_df.iloc[0]
        n_knots = len(month_df)
        knots_t = np.arange(0, n_knots + 1) * (STEP_HOURS * 3600.0)
        knots = {c: np.concatenate([[float(knot_prev[c])], month_df[c].to_numpy(float)]) for c in SEC_COLS}
        total_sec = n_knots * STEP_HOURS * 3600

        use_parquet = SECONDS_FORMAT == "parquet"
        if use_parquet:
            try:
                import pyarrow as pa, pyarrow.parquet as pq
            except ImportError:
                use_parquet = False
                print("pyarrow not available -> writing csv.gz", flush=True)
        file_path = OUT_LOCAL / (f"ADM1_seconds_{period}.parquet" if use_parquet else f"ADM1_seconds_{period}.csv.gz")

        writer = None
        fh = None
        try:
            for day0 in range(0, total_sec, 86400):
                t = np.arange(day0 + 1, min(day0 + 86400, total_sec) + 1, dtype=float)   # seconds from month start (1 s ...)
                chunk = {"datetime": month_start + pd.to_timedelta(t, unit="s")}
                for c in SEC_COLS:
                    v = np.interp(t, knots_t, knots[c]).astype(np.float32)
                    if c.startswith("Q_"):
                        chunk[c.replace("_m3_d", "_m3_s")] = (v / 86400.0).astype(np.float32)
                    else:
                        chunk[c] = v
                cdf = pd.DataFrame(chunk)
                if use_parquet:
                    table = pa.Table.from_pandas(cdf, preserve_index=False)
                    if writer is None:
                        writer = pq.ParquetWriter(file_path, table.schema, compression="zstd")
                    writer.write_table(table)
                else:
                    if fh is None:
                        fh = gzip.open(file_path, "wt")
                        cdf.to_csv(fh, index=False, float_format="%.6g")
                    else:
                        cdf.to_csv(fh, index=False, header=False, float_format="%.6g")
        finally:
            if writer is not None:
                writer.close()
            if fh is not None:
                fh.close()
        save_copy(file_path)
        print("Second data saved:", file_path.name, f"({total_sec} rows)", flush=True)

    last_record = month_df.iloc[-1]

# ============================================================
# 8. Save complete time series and monthly volumes
# ============================================================

series = pd.concat(all_frames, ignore_index=True)
timeseries_file = OUT_LOCAL / "ADM1_timeseries.csv"
series.to_csv(timeseries_file, index=False)
save_copy(timeseries_file)

for c in ["Q_gas_wet_m3_d_20C", "Q_gas_dry_m3_d_20C", "Q_CH4_m3_d", "Q_CO2_m3_d", "Q_H2_m3_d", "Q_H2O_vapour_m3_d"]:
    series[c + "_vol"] = series[c] * dt_day   # end-of-interval rate x interval length = interval volume [m3]
series["H2O_kg_d_vol"] = series["H2O_kg_d"] * dt_day

monthly = series.groupby("month", as_index=False).agg(
    gas_wet_m3=("Q_gas_wet_m3_d_20C_vol", "sum"),
    gas_dry_m3=("Q_gas_dry_m3_d_20C_vol", "sum"),
    CH4_m3=("Q_CH4_m3_d_vol", "sum"),
    CO2_m3=("Q_CO2_m3_d_vol", "sum"),
    H2_m3=("Q_H2_m3_d_vol", "sum"),
    H2O_kg=("H2O_kg_d_vol", "sum"),
    CH4_dry_pct=("CH4_dry_pct", "mean"),
    CO2_dry_pct=("CO2_dry_pct", "mean"),
    H2_dry_ppm=("H2_dry_ppm", "mean"),
    H2O_wet_pct=("H2O_wet_pct", "mean"),
    pH_min=("pH", "min"),
    pH_end=("pH", "last"),
    VFA_end=("VFA_kgCOD_m3", "last"),
    TAN_end=("TAN_gN_L", "last"),
    NH3_max=("NH3_mgN_L", "max"),
    NH3_gas_ppmv_mean=("NH3_gas_ppmv_eq", "mean"),
    NH3_gas_kg=("NH3_gas_g_d_eq", lambda x: x.sum() * dt_day / 1000.0),
    TAN_out_kgN=("TAN_out_kgN_d", lambda x: x.sum() * dt_day),
    H2S_gas_ppmv_mean=("H2S_gas_ppmv_eq", "mean"),
    H2S_gas_ppmv_max=("H2S_gas_ppmv_eq", "max"),
    H2S_gas_kg=("H2S_gas_kg_d_eq", lambda x: x.sum() * dt_day),
    sulfide_out_kgS=("sulfide_out_kgS_d", lambda x: x.sum() * dt_day),
)
monthly_file = OUT_LOCAL / "ADM1_monthly.csv"
monthly.to_csv(monthly_file, index=False)
save_copy(monthly_file)
print("\nMONTHLY SUMMARY")
print(monthly.round(2).to_string(index=False))

# ============================================================
# 9. Plots: gas, composition, pH
# ============================================================

fig, axs = plt.subplots(3, 1, figsize=(12, 10), sharex=True)
axs[0].plot(series["interval_end"], series["Q_gas_wet_m3_d_20C"], label="biogas (wet)")
axs[0].plot(series["interval_end"], series["Q_CH4_m3_d"], label="CH4")
axs[0].set_ylabel("m3/day (20 C ref)"); axs[0].legend()
axs[1].plot(series["interval_end"], series["CH4_dry_pct"], label="CH4 dry %")
axs[1].plot(series["interval_end"], series["CO2_dry_pct"], label="CO2 dry %")
axs[1].plot(series["interval_end"], series["H2_dry_ppm"] / 1e4, label="H2 dry %")
axs[1].plot(series["interval_end"], series["H2O_wet_pct"], label="H2O wet %")
axs[1].set_ylabel("vol %"); axs[1].legend()
axs[2].plot(series["interval_end"], series["pH"]); axs[2].set_ylabel("pH"); axs[2].set_xlabel("Date")
for ax in axs:
    ax.grid(alpha=0.3)
fig.tight_layout()
png = OUT_LOCAL / "ADM1_gas_composition_pH.png"
fig.savefig(png, dpi=200)
save_copy(png)
plt.show()

# ============================================================
# 9b. H2S / NH3 concentration plots (ppmv, dry gas basis) — separate figure
# ============================================================

fig2, ax2 = plt.subplots(2, 1, figsize=(12, 8), sharex=True)

# Monthly mean and max for overlay (by interval_end)
_m = series[["H2S_gas_ppmv_eq", "NH3_gas_ppmv_eq"]].copy()
_m.index = series["interval_end"] - pd.Timedelta(seconds=1)      # keep the 00:00 interval end from rolling into the next month
_mon_mean = _m.resample("MS").mean()
_mon_max = _m.resample("MS").max()

ax2[0].plot(series["interval_end"], series["H2S_gas_ppmv_eq"], color="tab:red", lw=1.0, label="H2S (hourly, equilibrium estimate)")
ax2[0].step(_mon_mean.index, _mon_mean["H2S_gas_ppmv_eq"], where="post", color="k", lw=1.2, ls="--", label="monthly mean")
ax2[0].set_ylabel("H2S [ppmv, dry]")
ax2[0].set_ylim(bottom=0)
ax2[0].set_title(f"H2S in digester off-gas  (F_S_PROTEIN={F_S_PROTEIN}, SO4_in={SO4_IN_mgS_L} mg S/L, Fe={FE_DOSE_mol_m3} mol/m3)")
ax2[0].legend(loc="upper right")

ax2[1].plot(series["interval_end"], series["NH3_gas_ppmv_eq"], color="tab:blue", lw=1.0, label="NH3 (hourly, Henry equilibrium)")
ax2[1].step(_mon_mean.index, _mon_mean["NH3_gas_ppmv_eq"], where="post", color="k", lw=1.2, ls="--", label="monthly mean")
ax2[1].set_ylabel("NH3 [ppmv, dry]")
ax2[1].set_ylim(bottom=0)
ax2[1].set_title("NH3 in digester off-gas (equilibrium with free NH3 in liquid; upper-bound estimate)")
ax2[1].set_xlabel("Date")
ax2[1].legend(loc="upper right")

for ax in ax2:
    ax.grid(alpha=0.3)
    # annotate monthly maxima
for ax, col in zip(ax2, ["H2S_gas_ppmv_eq", "NH3_gas_ppmv_eq"]):
    for ts, val in _mon_max[col].items():
        if np.isfinite(val):
            ax.annotate(f"{val:.0f}", (ts + pd.Timedelta(days=15), val), textcoords="offset points", xytext=(0, 4),
                        ha="center", fontsize=7, color="dimgray")

fig2.tight_layout()
png2 = OUT_LOCAL / "ADM1_H2S_NH3_ppm.png"
fig2.savefig(png2, dpi=200)
save_copy(png2)
plt.show()

# Monthly ppm summary (mean / max) saved separately
ppm_monthly = pd.DataFrame({
    "H2S_ppmv_mean": _mon_mean["H2S_gas_ppmv_eq"], "H2S_ppmv_max": _mon_max["H2S_gas_ppmv_eq"],
    "NH3_ppmv_mean": _mon_mean["NH3_gas_ppmv_eq"], "NH3_ppmv_max": _mon_max["NH3_gas_ppmv_eq"],
}).round(1)
ppm_monthly.index = ppm_monthly.index.strftime("%Y-%m")
ppm_file = OUT_LOCAL / "ADM1_H2S_NH3_ppm_monthly.csv"
ppm_monthly.to_csv(ppm_file)
save_copy(ppm_file)
print("\nH2S / NH3 ppmv (monthly mean / max)")
print(ppm_monthly.to_string())

print("\n" + "=" * 60)
print("SIMULATION FINISHED")
print("Output:", OUT_DRIVE)
print("Total ADM1 steps:", len(series))
print("=" * 60)
if RUN_MODE == "test":
    print("Test completed. Set RUN_MODE='full' and rerun for 12 months.")
